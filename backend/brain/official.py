"""Official Caritas facility list (Caritas_Einrichtungen.xlsx) -> extra rows for the services table.

The list has 980 facilities with type, name and address, but no phone, email, hours or description.
This module keeps only the types a person looking for help would be routed to, and gives each type
an English/German description, everyday keywords and its online-counselling topics, so semantic search
can match "my mother needs care" to a "Sozialstation".

Left out on purpose: day care, schools and training, administration, hospitals, and women's shelters
(Frauenhaus): their locations are kept confidential for safety; violence goes to the Hilfetelefon instead.
"""
from pathlib import Path

import openpyxl

COLS = {"type": "Einrichtungsart", "name1": "Einrichtungsname 1", "name2": "Einrichtungsname 2",
        "street": "Straße", "postcode": "PLZ", "city": "Ort", "district": "Ortsteil"}

# (type prefix as in the file) -> (category EN, short description DE, topic ids, keywords EN, keywords DE)
TYPES = [
    # ---- general help / first contact
    ("Allgemeine Sozialberatung", "General social counselling", "Allgemeine Sozialberatung", ["OT01"],
     "general problems; don't know who to ask; money worries; forms and applications; letters from authorities; benefits",
     "Probleme; weiß nicht weiter; Geldsorgen; Anträge; Formulare; Behördenbriefe; Sozialleistungen; Beratung"),
    ("Orts-, Kreis-, Bezirks-Caritasverband", "Local Caritas office (first contact)", "Caritasverband vor Ort",
     ["OT01"], "Caritas office; first contact; who can help me; general enquiry",
     "Caritas vor Ort; Ansprechpartner; erste Anlaufstelle; allgemeine Anfrage"),
    ("Fachdienst Gemeindecaritas", "Caritas community service", "Gemeindecaritas", ["OT01"],
     "community help; neighbourhood help; volunteers; local support", "Gemeinde; Nachbarschaftshilfe; Ehrenamt; Hilfe vor Ort"),
    ("Diözesan-Caritasverband", "Diocesan Caritas association (head office)", "Diözesan-Caritasverband", ["OT01"],
     "Caritas head office; general enquiry", "Caritas Zentrale; allgemeine Anfrage"),
    ("Diözesan-, Bezirks- oder Ortsverband SkF", "Catholic women's social service (SkF)", "Sozialdienst katholischer Frauen",
     ["OT04", "OT15", "OT01"], "women; mothers; pregnancy; family problems; legal guardianship",
     "Frauen; Mütter; Schwangerschaft; Familienprobleme; Betreuung"),
    ("Beratung/Betreuung im eigenen Wohnraum", "Support in your own home", "Betreuung zu Hause", ["OT03", "OT01"],
     "support at home; help living independently", "Unterstützung zu Hause; selbstständig wohnen"),
    ("Mehrgenerationenhaus", "Multi-generation house (meeting place)", "Mehrgenerationenhaus", ["OT11", "OT01"],
     "meeting place; lonely; contact with others; activities for all ages", "Begegnung; einsam; Kontakte; Treffpunkt"),
    ("Quartiersarbeit/Gemeinwesensarbeit", "Neighbourhood work", "Quartiersarbeit", ["OT01"],
     "neighbourhood; community; local projects", "Stadtteil; Nachbarschaft; Gemeinwesen"),
    ("Freiwilligen-Zentrum", "Volunteer centre", "Freiwilligen-Zentrum", ["OT01"],
     "volunteering; help others; volunteer work", "Ehrenamt; freiwillig helfen; engagieren"),
    # ---- everyday needs
    ("Bahnhofsmission", "Bahnhofsmission (help at the station)", "Bahnhofsmission", ["OT01", "HX01"],
     "help at the train station; stranded; travelling alone; nowhere to go; a warm drink",
     "Hilfe am Bahnhof; gestrandet; allein unterwegs; nicht wissen wohin"),
    ("Suppenküche/Tafel", "Food bank / soup kitchen", "Tafel / Suppenküche", ["OT01"],
     "food; hungry; no money for food; free meals; food bank", "Essen; Hunger; kein Geld für Essen; Lebensmittel; Tafel"),
    ("Kleiderkammer/Möbellager", "Clothing and furniture store (low cost)", "Kleiderkammer / Sozialkaufhaus", ["OT01"],
     "clothes; furniture; cheap second-hand; no money for clothes", "Kleidung; Möbel; günstig; gebraucht; Sozialkaufhaus"),
    # ---- money
    ("Schuldner-/Insolvenzberatungsstelle", "Debt and insolvency counselling", "Schuldnerberatung", ["OT14"],
     "debt; bills; collection letters; insolvency; can't pay", "Schulden; Rechnungen; Mahnungen; Inkasso; Insolvenz"),
    # ---- addiction
    ("Beratungs- oder Behandlungsstelle für Suchtkranke", "Addiction counselling", "Suchtberatung", ["OT17"],
     "alcohol; drinking too much; drugs; gambling; addiction; relatives of addicts",
     "Alkohol; trinken; Drogen; Glücksspiel; Sucht; Angehörige"),
    # ---- migration
    ("Migrationsdienst: Flüchtlingssozial- und Asylverfahrensberatung", "Refugee and asylum counselling",
     "Flüchtlings- und Asylberatung", ["OT12"], "refugee; asylum; papers; residence; new in Germany",
     "Flüchtling; Asyl; Papiere; Aufenthalt; neu in Deutschland"),
    ("Migrationsdienst: Landes- und kommunal finanzierte Integrationsberatung", "Integration counselling",
     "Integrationsberatung", ["OT12"], "integration; language course; settling in; authorities",
     "Integration; Sprachkurs; ankommen; Behörden"),
    ("Migrationsdienst: Migrationsberatung für Erwachsene", "Migration counselling for adults (MBE)",
     "Migrationsberatung für Erwachsene", ["OT12"], "migrant; immigration; work; language; settling in",
     "Zuwanderer; Migration; Arbeit; Sprache; ankommen"),
    ("Migrationsdienst: Sozialpädagogischer Dienst für ausländische Kinder", "Migration service for children and youth",
     "Migrationsdienst für Kinder und Jugendliche", ["OT12", "OT08"], "migrant children; young people; school; integration",
     "Kinder; Jugendliche; Migration; Schule; Integration"),
    # ---- pregnancy and families
    ("Schwangerschaftsberatungsstelle", "Pregnancy counselling", "Schwangerschaftsberatung", ["OT15"],
     "pregnant; pregnancy; baby; unplanned pregnancy; financial help for mothers",
     "schwanger; Schwangerschaft; Baby; ungeplant; finanzielle Hilfe"),
    ("Gemeinsame Wohnform für Schwangere/Mütter und Väter", "Supported housing for pregnant women and young parents",
     "Wohnform für Schwangere und junge Eltern", ["OT15", "OT04"], "pregnant and nowhere to live; young mother; housing with baby",
     "schwanger und keine Wohnung; junge Mutter; Wohnen mit Baby"),
    ("Beratungsstelle für Frühe Hilfen", "Early help for families with babies", "Frühe Hilfen", ["OT04", "OT15"],
     "new baby; overwhelmed parents; support after birth", "Baby; überforderte Eltern; Unterstützung nach der Geburt"),
    ("Erziehungsberatungsstelle", "Parenting and family counselling", "Erziehungsberatung", ["OT04", "OT08"],
     "parenting; problems with my child; teenager; separation; family conflict",
     "Erziehung; Probleme mit dem Kind; Pubertät; Trennung; Familienstreit"),
    ("Beratungsstelle für Ehe-, Familien- und Lebensfragen", "Marriage, family and life counselling",
     "Ehe-, Familien- und Lebensberatung", ["OT04", "OT01"], "relationship problems; marriage; separation; life crisis",
     "Beziehungsprobleme; Ehe; Trennung; Lebenskrise"),
    ("Familienzentrum", "Family centre", "Familienzentrum", ["OT04"], "family; parents; children; courses; meeting",
     "Familie; Eltern; Kinder; Kurse; Treffpunkt"),
    ("Sozialpädagogische Familienhilfe", "Family support at home", "Sozialpädagogische Familienhilfe", ["OT04"],
     "family support at home; overwhelmed with everyday life and children", "Familienhilfe; Alltag mit Kindern; überfordert"),
    ("Ambulante Erziehungshilfe", "Outpatient parenting support", "Ambulante Erziehungshilfe", ["OT04", "OT08"],
     "help with raising children; support for parents", "Erziehungshilfe; Unterstützung für Eltern"),
    ("Familienpflegestation", "Family care service (when a parent is ill)", "Familienpflege", ["OT04"],
     "parent ill; household and children need help", "Elternteil krank; Haushalt und Kinder versorgen"),
    # ---- children and youth
    ("Tagesstätte für Jugendliche", "Youth centre", "Jugendtreff", ["OT08"],
     "youth; teenagers; meeting place; free time", "Jugendliche; Treffpunkt; Freizeit"),
    ("Schulsozialarbeit", "School social work", "Schulsozialarbeit", ["OT08"],
     "problems at school; bullying; pupils", "Probleme in der Schule; Mobbing; Schüler"),
    ("Mädchen- und Frauentreff", "Girls' and women's meeting place", "Mädchen- und Frauentreff", ["OT08", "OT04"],
     "girls; women; meeting place", "Mädchen; Frauen; Treffpunkt"),
    ("Heimerziehung/sonstige betreute Wohnform", "Supported housing for children and youth", "Heimerziehung / betreutes Wohnen",
     ["OT08"], "young people who cannot live at home", "Jugendliche, die nicht zu Hause wohnen können"),
    ("Intensive Sozialpädagogische Einzelbetreuung", "Intensive one-to-one youth support", "Intensive Einzelbetreuung",
     ["OT08"], "young person in difficulties; individual support", "Jugendliche in Schwierigkeiten; Einzelbetreuung"),
    # ---- older people and care
    ("Sozialstation/ambulanter Pflegedienst", "Home care service (Sozialstation)", "Ambulanter Pflegedienst", ["OT11"],
     "care at home; nursing at home; elderly parent needs care; help with washing and medication",
     "Pflege zu Hause; ambulante Pflege; pflegebedürftig; Hilfe beim Waschen; Medikamente"),
    ("Sozialstationen/ambulante Pflegedienste mit angeschlossener Familienpflege", "Home care and family care service",
     "Ambulanter Pflegedienst mit Familienpflege", ["OT11", "OT04"],
     "care at home; family care; elderly; ill parent", "Pflege zu Hause; Familienpflege; pflegebedürftig"),
    ("Tagespflegeeinrichtung", "Day care for older people", "Tagespflege", ["OT11"],
     "day care for elderly; relief for family carers; dementia during the day",
     "Tagespflege; Entlastung für Angehörige; Demenz tagsüber"),
    ("Vollstationäre Altenpflegeeinrichtung", "Care home for older people", "Pflegeheim", ["OT11"],
     "care home; nursing home; elderly can no longer live alone", "Pflegeheim; Altenheim; kann nicht mehr allein leben"),
    ("Kurzzeitpflegeeinrichtung", "Short-term care", "Kurzzeitpflege", ["OT11"],
     "short-term care; after hospital; carer on holiday", "Kurzzeitpflege; nach dem Krankenhaus; Urlaub pflegender Angehöriger"),
    ("Altenwohnanlage/Betreutes Wohnen", "Assisted living for older people", "Betreutes Wohnen", ["OT11"],
     "assisted living; senior apartment", "betreutes Wohnen; Seniorenwohnung"),
    ("Altenwohnheim", "Residential home for older people", "Altenwohnheim", ["OT11"],
     "senior residence; old people's home", "Altenwohnheim; Seniorenheim"),
    ("Altentagesstätte, Altenbegegnungsstätte", "Senior meeting centre", "Seniorentreff", ["OT11"],
     "elderly; lonely; meeting place; activities for seniors", "Senioren; einsam; Treffpunkt; Freizeit"),
    ("Mahlzeitendienst, Stationärer Mittagstisch, Essen auf Rädern", "Meals on wheels / lunch service",
     "Essen auf Rädern / Mittagstisch", ["OT11", "OT01"], "meals delivered; lunch; cannot cook anymore",
     "Essen auf Rädern; Mittagstisch; kann nicht mehr kochen"),
    ("Dauerpflegeeinrichtung nach SGB XI in einer Einrichtung der Behindertenhilfe",
     "Long-term care for people with disabilities", "Dauerpflege in der Behindertenhilfe", ["OT03", "OT11"],
     "long-term care; disability; care home", "Dauerpflege; Behinderung; Pflege"),
    ("Seniorenbetreuung für behinderte Menschen", "Support for older people with disabilities",
     "Seniorenbetreuung für Menschen mit Behinderung", ["OT03", "OT11"], "older people with disabilities",
     "ältere Menschen mit Behinderung"),
    # ---- serious illness
    ("Beratungs-/Betreuungsstelle für schwerstkranke Menschen", "Counselling for seriously ill people",
     "Beratung für schwerstkranke Menschen", ["OT06", "OT05"], "cancer; serious illness; HIV; dying; palliative",
     "Krebs; schwer krank; HIV; Sterben; Palliativ"),
    # ---- disability and mental health
    ("Beratungs-/Betreuungsstelle für behinderte Menschen", "Counselling for people with disabilities",
     "Beratung für Menschen mit Behinderung", ["OT03"], "disability; support; family relief",
     "Behinderung; Unterstützung; Entlastung der Familie"),
    ("Ambulanter sozialpsychiatrischer Dienst", "Social psychiatric service", "Sozialpsychiatrischer Dienst", ["OT03"],
     "mental health; depression; psychological problems; anxiety", "psychische Probleme; Depression; Angst; seelische Krise"),
    ("Tagesstätte für erwachsene psychisch kranke", "Day centre for people with mental illness",
     "Tagesstätte für psychisch kranke Menschen", ["OT03"], "mental illness; day structure; contact", "psychisch krank; Tagesstruktur; Kontakte"),
    ("Ambulant betreutes Wohnen für psychisch kranke", "Supported living for people with mental illness",
     "Betreutes Wohnen für psychisch kranke Menschen", ["OT03"], "supported living; mental illness",
     "betreutes Wohnen; psychische Erkrankung"),
    ("Ambulant betreutes Wohnen für behinderte", "Supported living for people with disabilities",
     "Betreutes Wohnen für Menschen mit Behinderung", ["OT03"], "supported living; disability", "betreutes Wohnen; Behinderung"),
    ("Übergangsheim für psychisch kranke", "Transitional home for people with mental illness", "Übergangsheim",
     ["OT03"], "after psychiatric hospital; transitional housing", "nach der Psychiatrie; Übergangswohnen"),
    ("Heim/Wohnheim für psychisch kranke", "Residential home for people with mental illness", "Wohnheim für psychisch kranke Menschen",
     ["OT03"], "residential home; mental illness", "Wohnheim; psychische Erkrankung"),
    ("Heim/Wohnheim für lernbehinderte/geistig behinderte", "Residential home for people with intellectual disabilities",
     "Wohnheim für Menschen mit geistiger Behinderung", ["OT03"], "intellectual disability; residential home",
     "geistige Behinderung; Wohnheim"),
    ("Außenwohngruppe für lernbehinderte", "Supported group home for people with learning disabilities",
     "Außenwohngruppe", ["OT03"], "group home; learning disability", "Wohngruppe; Lernbehinderung"),
    ("Werkstätte für behinderte Menschen", "Sheltered workshop for people with disabilities", "Werkstatt für Menschen mit Behinderung",
     ["OT03"], "work for people with disabilities; workshop", "Arbeit für Menschen mit Behinderung; Werkstatt"),
    ("Förderstätte für erwachsene behinderte Menschen", "Day support for adults with severe disabilities", "Förderstätte",
     ["OT03"], "adults with severe disabilities; day support", "Erwachsene mit schwerer Behinderung; Tagesförderung"),
    ("Einrichtung zur Früherkennung, Frühbehandlung und Frühförderung", "Early intervention for children",
     "Frühförderung", ["OT03", "OT08"], "child development delay; early support; toddler", "Entwicklungsverzögerung; Frühförderung; Kleinkind"),
    ("Sonstige Beratungs-/Betreuungsstelle der Gesundheitshilfe", "Health counselling", "Beratungsstelle der Gesundheitshilfe",
     ["OT03"], "health; illness; counselling", "Gesundheit; Krankheit; Beratung"),
    ("Medizinische und/ oder berufliche Rehabilitationseinrichtung für Menschen mit psychischer",
     "Rehabilitation for people with mental illness", "Rehabilitation bei psychischer Erkrankung", ["OT03"],
     "rehabilitation; back to work; mental illness", "Rehabilitation; zurück in den Beruf; psychische Erkrankung"),
    ("Betreuungsverein", "Legal guardianship association", "Betreuungsverein", ["OT13"],
     "legal guardian; power of attorney; advance directive", "rechtliche Betreuung; Vorsorgevollmacht; Patientenverfügung"),
    # ---- housing, offending
    ("Beratungsstelle für Personen in bes. soz. Schwierigkeiten/Wohnungslose", "Counselling for homeless people",
     "Beratung für Wohnungslose", ["HX01", "OT01"], "homeless; lost my flat; eviction; nowhere to sleep",
     "wohnungslos; Wohnung verloren; Räumung; nirgends schlafen"),
    ("Stationäre Einrichtung f. Personen in bes. soz. Schwierigkeiten", "Residential help for people in social difficulties",
     "Stationäre Hilfe in sozialen Schwierigkeiten", ["HX01"], "homeless; residential help; social difficulties",
     "wohnungslos; stationäre Hilfe; soziale Schwierigkeiten"),
    ("Tagesstätte für Personen in bes. soz. Schwierigkeiten", "Day centre for people in social difficulties",
     "Tagesstätte in sozialen Schwierigkeiten", ["HX01", "OT01"], "day centre; homeless; warm place; meal",
     "Tagesstätte; wohnungslos; Aufenthalt; Essen"),
    ("Wohngruppe/Betreutes Wohnen für Wohnungslose", "Supported housing for homeless people", "Betreutes Wohnen für Wohnungslose",
     ["HX01"], "homeless; supported housing", "wohnungslos; betreutes Wohnen"),
    ("Beratungsstelle für Straffällige und Haftentlassene", "Counselling for offenders and people leaving prison",
     "Straffälligenhilfe", ["OT16"], "prison; released from prison; criminal record", "Gefängnis; Haftentlassung; Straffälligkeit"),
    # ---- crisis
    ("Telefonseelsorgestelle", "Telefonseelsorge (crisis line office)", "Telefonseelsorge", ["OT19", "OT18"],
     "someone to talk to; crisis; lonely; desperate", "jemand zum Reden; Krise; einsam; verzweifelt"),
    # ---- recovery
    ("Vorsorge- u. Rehabilitationseinrichtung für Erwachsene", "Health resort / rehabilitation for adults",
     "Vorsorge- und Rehabilitationseinrichtung", ["OT10"], "rehabilitation; recovery stay; exhausted mother",
     "Kur; Rehabilitation; Erholung; erschöpft"),
]

EXCLUDED_FOR_SAFETY = ("Frauenhaus",)


def _type_info(facility_type: str):
    for prefix, *info in TYPES:
        if facility_type.startswith(prefix):
            return info
    return None


def load(path, region_of_city, geo=None, region_of_county=None):
    """Return service rows (same columns as the team's services) for the help-relevant facility types."""
    path = Path(path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    headers = [str(h).strip() if h is not None else None for h in next(rows)]
    missing = [h for h in COLS.values() if h not in headers]
    if missing:
        raise KeyError(f"{path.name} is missing columns: {missing}")
    idx = {h: i for i, h in enumerate(headers) if h}

    raw = []
    for row in rows:
        r = {k: (str(row[idx[h]]).strip() if row[idx[h]] is not None else "") for k, h in COLS.items()}
        if r["type"] and not any(word in r["type"] for word in EXCLUDED_FOR_SAFETY):
            raw.append(r)

    # Region per town: district towns are known by name; every other town takes the region of the facility
    # with the numerically nearest postcode (German postcodes are assigned by area, so this approximates
    # "nearest district town"; good enough to prefer nearby places, not an official district lookup).
    from .textutil import normalize

    geo = geo or {}
    for r in raw:
        g = geo.get(normalize(r["city"]))
        r["region"] = (region_of_county(g["county"]) if g and region_of_county else None) or region_of_city(r["city"])
    known = [(int(r["postcode"]), r["region"]) for r in raw if r["region"] and r["postcode"].isdigit()]
    for r in raw:
        if not r["region"] and r["postcode"].isdigit() and known:
            plz = int(r["postcode"])
            r["region"] = min(known, key=lambda k: abs(k[0] - plz))[1]

    services, links = [], []
    for n, r in enumerate(raw, 1):
        info = _type_info(r["type"])
        if info is None:
            continue  # day care, schools, administration, hospitals: not routed to by this assistant
        category, label_de, topics, kw_en, kw_de = info
        n1, n2 = r["name1"], r["name2"]
        if not n2 or n2 in n1:
            name = n1
        elif n1.endswith(("-", " und", " für", " der", " des", " im", " am")) or n2[:1].islower():
            name = f"{n1} {n2}"               # one name split over two columns: "Schuldner- und" + "Insolvenzberatung"
        else:
            name = f"{n1} – {n2}"
        city = r["city"]
        street = r["street"] + (f" ({r['district']})" if r["district"] else "")
        sid = f"CE{n:04d}"
        services.append({
            "id": sid, "kind": "facility", "region": r["region"], "category": category, "name": name,
            "provider": "Caritas", "relationship": "Official Caritas facility list", "city": city,
            "postcode": r["postcode"], "address": street.strip() or None, "coverage": f"{label_de} in {city}",
            "phone": None, "email": None, "website": None, "contact_mode": None, "hours": None,
            "access_rule": None, "cost_note": None, "keywords_en": kw_en, "keywords_de": kw_de,
            "confidence_note": "From the official Caritas facility list: name and address only, no phone number.",
            "source_url": None, "last_checked": None, "origin": f"Caritas_Einrichtungen.xlsx ({r['type']})",
            "search_text": " | ".join(p for p in (name, category, label_de, r["type"], city,
                                                   f"Keywords: {kw_en}", f"Stichworte: {kw_de}") if p),
        })
        links += [(sid, t) for t in topics]
    return services, links
